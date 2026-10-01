"""
The Production canvas, as a Word guide with pictures.

    python3 scripts/make-canvas-guide.py

Writes ~/Documents/Film Engine — Production Canvas Guide.docx from the
pictures in docs/guide/production-canvas/. The source lives here, beside the
code it describes, because the guide before this one was generated from a
temporary folder and could not be rebuilt once that folder was cleaned up.
tests/canvas-guide.test.js holds the text to the page: every feature, every
right-click item, header control, queue bucket, node state and agent tool.

Needs python-docx (pip3 install python-docx). Nothing in the backend uses it.
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'guide', 'production-canvas')
OUT = os.path.expanduser('~/Documents/Film Engine — Production Canvas Guide.docx')

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
WIDTH = Inches(6.7)

styles = doc.styles
styles['Normal'].font.name = 'Calibri'
styles['Normal'].font.size = Pt(10.5)
for name, size in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
    styles[name].font.name = 'Calibri'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def p(text, bold_lead=None, style=None, italic=False):
    para = doc.add_paragraph(style=style)
    if bold_lead:
        r = para.add_run(bold_lead); r.bold = True
    r = para.add_run(text); r.italic = italic
    para.paragraph_format.space_after = Pt(6)
    return para


def bullets(items):
    for it in items:
        if isinstance(it, tuple):
            p(it[1], bold_lead=it[0], style='List Bullet')
        else:
            p(it, style='List Bullet')


def pic(name, caption, width=None):
    path = os.path.join(IMG, name + '.jpg')
    doc.add_picture(path, width=width or WIDTH)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x60, 0x66, 0x70)


def pics_side(names, caption, each=Inches(2.15)):
    t = doc.add_table(rows=1, cols=len(names))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, n in zip(t.rows[0].cells, names):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(os.path.join(IMG, n + '.jpg'), width=each)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x60, 0x66, 0x70)


def table(header, rows, widths=None, fill='1F2A44'):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.size = Pt(9.5)
        shade(cell, fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(v); run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    doc.add_paragraph()
    return t


def node_card(what, how, settings):
    p(what, bold_lead='What it is. ')
    p(how, bold_lead='How it works. ')
    if settings:
        p('', bold_lead='Settings (in its drawer):')
        bullets(settings)


# ── Cover ─────────────────────────────────────────────────────────────────
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run('Film Engine'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x60, 0x66, 0x70)
doc.add_heading('The Production Canvas', 0)
p('How the graph works, what every node does, and every setting you can change — including what the canvas now '
  'shows about the work itself: what is running and how far along, what is out of date, what a clip still needs, '
  'and how any version was made. The pictures are taken from the test film "Last Call (E2E)": one scene, three shots '
  '(1A, 1B, 1C), two sequences, a linked frame, a score and an ambient bed. Some of its states (a frame behind, a job '
  'running, a held shot) were set up on a copy of the project so each one could be photographed.', italic=True)
pic('00-full', 'The whole Production page: header, side panel, the canvas with its nodes, the queue, and the playback bar.')

doc.add_heading('Contents', level=2)
bullets(['1. Switching the canvas on', '2. The page at a glance', '3. How the canvas is laid out',
         '4. Is it up to date? Node states', '5. The nodes: Shot, Sequence, Linked frame, Video version, Sound, Audio version',
         '6. Live progress on the node', '7. Out-of-date on the graph, and "Run what changed"', '8. Run to here',
         '9. Queue and history strip — and what Cancel really does', '10. How was this made, Make another, Compare',
         '11. Hold a node', '12. Collapse a group', '13. Search to add, and patterns (double-click)',
         '14. Wiring: what connects to what', '15. Adding, duplicating and removing', '16. The header',
         '17. The playback bar', '18. What costs money, and what is free', '19. Asking Claude to do it',
         '20. Flows on the canvas', '21. The Score page', '22. Your own instruments',
         '23. An edit made in Premiere', '24. The final film and its sound', '25. Sending the film to Premiere',
         '26. Delivery quality: the resolution you asked for', '27. Upscaling a clip', '28. Your recorded dialogue in a clip',
         '29. Backups', '30. Setup: providers, models and keys', '31. A video model\'s own options',
         '32. Changing a project\'s stage', '33. Good to know', '34. Previs: how it works'])

# ── 1 ─────────────────────────────────────────────────────────────────────
doc.add_heading('1. Switching the canvas on', level=1)
p('The canvas is behind a setting. Open Setup (top right) → Settings and tick "Show production as one graph". '
  'With it on, the Production phase is this one page; the separate pages it replaces (Shot board, Video shots, '
  'Music, Music cues, Playback, Pipeline and Flows) leave the menu but stay reachable by address. '
  'The Score workspace stays in the menu either way, because a Music node opens it rather than replacing it.')
pic('32-setting', 'Settings → "Show production as one graph".', width=Inches(4.5))

# ── 2 ─────────────────────────────────────────────────────────────────────
doc.add_heading('2. The page at a glance', level=1)
table(['Area', 'Where', 'What it is for'], [
    ['Header', 'Top', 'Counts for the film, the state legend with "Only what\'s behind", scene filter chips, Tidy layout, Lock board, Run what changed, Run pending.'],
    ['Side panel', 'Left', 'An outline of every shot by scene (click to jump to it), and the palette you drag new nodes from.'],
    ['Canvas', 'Middle', 'The nodes and the wires between them. Drag the background to pan, scroll to zoom, double-click to add.'],
    ['Drawer', 'Right, when a node is selected', 'Everything about the selected node: whether it is up to date, how it was made, its settings, its versions, and its Generate button.'],
    ['Queue strip', 'Under the canvas', 'What is running, waiting, done today, waiting to be collected and failed — with Cancel, Collect and Re-run.'],
    ['Playback bar', 'Bottom', 'The film in running order, with a monitor, transport, and a lane per sound.'],
], widths=[Inches(1.2), Inches(1.6), Inches(3.9)])
pic('01-header', 'The header: counts, the legend, scene chips and the four header buttons.')
pics_side(['02-side', '03-furniture'], 'Left: the side panel — the shot outline and the "Add node" palette. '
          'Right: the canvas controls — zoom, Fit, and the colour legend for the wires.', each=Inches(2.6))

# ── 3 ─────────────────────────────────────────────────────────────────────
doc.add_heading('3. How the canvas is laid out', level=1)
p('Every node is a view of something that already exists in the project: a shot and its frames, a sequence, '
  'a sound cue, or a file a generation made. Nothing on the canvas is a separate copy — change it here and the '
  'board, the Score page and the exports all see the same change.')
doc.add_heading('Boxes', level=3)
bullets([
    ('A box per sequence. ', 'Each sequence gets a labelled box ("SC 1 · WE\'RE CLOSED") holding its shots, the sequence node, '
     'its video versions and its sounds.'),
    ('A box per scene for loose shots. ', 'Shots that belong to no sequence are grouped by scene ("SC 1 · DINER").'),
    ('Film order, left to right. ', 'Boxes run in the order the film plays, by each box\'s first shot. '
     'A new, empty sequence goes at the far right until you give it shots.'),
    ('Folding a box. ', 'The small arrow at a box\'s top right collapses it to one summary card (see section 12).'),
])
doc.add_heading('Moving around', level=3)
bullets([
    ('Pan: ', 'drag the empty background, or scroll sideways / shift-scroll.'),
    ('Zoom: ', 'scroll up and down (or pinch / ctrl-scroll), or use − / + in the bottom-left. "Fit" frames everything, collapsed cards included.'),
    ('Minimap: ', 'the small map bottom-right. Click anywhere on it to jump there.'),
    ('Scene chips: ', 'in the header, show one scene only, or All scenes.'),
    ('Move a node: ', 'drag it. A node you have dragged is pinned and stays put; "Tidy layout" re-places only the nodes nobody moved.'),
    ('Select: ', 'click a node to open its drawer; click the background, press Esc or the × to close it.'),
])
doc.add_heading('Ports and wire colours', level=3)
p('Each node has small round ports on its edges: inputs on the left, outputs on the right. A wire can only join ports of '
  'the same colour.')
table(['Colour', 'Carries', 'Example'], [
    ['Purple', 'image', 'A shot\'s frame into a sequence slot; a linked frame into a sequence.'],
    ['Yellow', 'video', 'A shot or sequence to its video versions.'],
    ['Green', 'audio', 'A sound to its audio versions.'],
    ['Blue / violet', 'scene', 'A shot or sequence into a sound, to give the sound its scene details.'],
    ['Cyan', 'plates', 'The reference pictures (characters, location, props) that travel with a shot.'],
], widths=[Inches(1.1), Inches(1.0), Inches(4.6)])
p('A hollow port means "nothing there yet". A dashed port is a linked-frame port. A wire takes the colour of the '
  'state of the node it leads into when that node is behind (red for redo, amber for waiting).')

# ── 4 ─────────────────────────────────────────────────────────────────────
doc.add_heading('4. Is it up to date? Node states', level=1)
p('Every node carries one of five states, read from the same impact report the rest of Film Engine uses — so the '
  'canvas never disagrees with the warnings on the storyboard. The state is the small pill at the node\'s top right and '
  'the colour of its border.')
pic('40-legend', 'The legend in the header: how many nodes are in each state, and the "Only what\'s behind" filter.', width=Inches(5.5))
table(['State', 'What it means', 'What to do'], [
    ['current', 'Made, and made from what is there now.', 'Nothing.'],
    ['redo', 'Out of date, and everything it is built from is current.', 'Regenerate it now — "Run what changed" will.'],
    ['waiting', 'Out of date only because something above it is (a sequence whose shot\'s frame is behind).', 'Wait: redo what it is built from first. Regenerating now would build on the old input and still be wrong.'],
    ['not made', 'Nothing has been generated yet.', 'Generate it when you are ready ("Run pending").'],
    ['untracked', 'Made outside the workflow (an upload, or made before tracking) — whether it is behind cannot be known.', 'Leave it, or regenerate it through Film Engine if you want it tracked.'],
], widths=[Inches(1.0), Inches(3.0), Inches(2.7)])
bullets([
    ('Why it is behind: ', 'hover the pill, or open the node — the drawer starts with the state, the reason ("What it was '
     'generated from has changed…") and the action.'),
    ('Badges on a shot: ', 'dialogue, lip-sync, sound effects and post roll up onto their shot as small badges, because they '
     'have no node of their own.'),
    ('Only what\'s behind: ', 'click it in the legend to hide everything that is current, leaving only the nodes that need '
     'work (and what they belong to). Click again to show everything.'),
    ('The film master: ', 'the whole-film assembly is not a node; when it is behind, the header says so and links to Conform.'),
])
pics_side(['41-node-redo', '45-drawer-impact'], 'Left: 1A is "redo" — its frame was made from a card that has since changed. '
          'Right: its drawer opens with the reason and what to do.', each=Inches(3.0))
pic('44-behind-only', '"Only what\'s behind" on: only the out-of-date work is left on the canvas.')

# ── 5 ─────────────────────────────────────────────────────────────────────
doc.add_heading('5. The nodes', level=1)

doc.add_heading('5.1 Shot', level=2)
pic('41-node-redo', 'A Shot node: code, framing, length, the selected frame, the version counter and its state pill.', width=Inches(2.6))
node_card(
    'One shot of the film: its storyboard frame and everything generated from it. The title is the shot code (1B); the '
    'line under it is the framing and length; "v1/1" is which frame version is selected out of how many exist. A red '
    '"no frame" pill means it still needs a picture. "↗ starts We\'re closed" means another sequence borrows this frame. '
    'A "HELD" tag and a dashed border mean batch runs skip it (section 11); a bar and "rendering · 62%" mean something is '
    'being made for it right now (section 6).',
    'The frame is generated from the shot\'s scene card plus its plates (the character, location and prop pictures that '
    'travel with it). Its "video" output leads to its video versions. Its "image" output can be wired into a sequence '
    'slot. Its "scene" output (top right) can be wired into a sound so the sound is written from this shot.',
    [('State: ', 'first in the drawer — redo / waiting / untracked with why and what to do.'),
     ('Frame: ', 'the selected picture, with a chip per version — click a chip to show that version everywhere. '
      'Regenerate makes a new version; Refine keeps the picture and changes one thing; 4 angles explores four cameras; '
      '+ Upload puts your own picture in; Anchor makes this frame the one later shots are shot from; Mark up opens the '
      'full-size viewer with drawing tools.'),
     ('How was v1 made: ', 'the recipe of the selected frame (section 10), with "Make another like this".'),
     ('From Previs: ', 'the directing decisions locked in Previs and their state. Changing them happens in Previs.'),
     ('Plates & details: ', 'the reference pictures this shot sends. Buttons to add an anchor frame, markup notes, a Previs '
      'angle, a Style book entry, or edit the card.'),
     ('Lighting: ', 'editing the card offers a lighting technique (Rembrandt, Loop, Window, Silhouette and thirteen more) '
      'and which side the key comes from, over the card\'s mood. A shot with none takes its location\'s; the film\'s '
      'general look stays the mood board\'s style. Staged in Previs and applied, it is written here (section 34).'),
     ('Video prompt — what is sent: ', 'the exact text the video model will receive, "By part" or "One text". Your edit '
      'replaces the composed prompt for this generation ("reset" undoes it).'),
     ('Video generation: ', 'Generator, Quality (Draft / Production / Hero), Duration, Seed.'),
     ('Dialogue: ', '"Generate dialogue" makes the take of every line on the card, in each character\'s cast voice.'),
     ('Video versions: ', 'every clip made for this shot — Open, Select to play it, Unselect, and Compare with selected.'),
     ('Footer: ', '"Preview request · free" shows the full request with nothing spent; "Generate video · $x" spends, '
      'after the confirmation.')])
pics_side(['20-drawer-shot-a', '21-drawer-shot-b', '22-drawer-shot-c'],
          'The Shot drawer, top to bottom: state, frame and recipe; Previs and plates; the video prompt, generation and versions.')

doc.add_heading('5.2 Sequence', level=2)
pic('11-node-sequence', 'A Sequence node: its shots in slots, the joins, its outputs, and its state.', width=Inches(2.8))
node_card(
    'Several shots made as one continuous move. The node lists its shots in order (1 · 1B, 2 · 1C). The rows on the left '
    'are input slots: "start frame" (dashed) for a borrowed opening frame, one slot per shot, "+ shot" to add one, and '
    '"end frame" for a borrowed closing frame. On the right: video (its clips), scene (to a sound), and end frame (to lend '
    'its last frame to another sequence). A sequence reads "waiting" while one of its shots\' frames is behind — its clip '
    'should be made after the frame is redone.',
    'A sequence travels between frames you have already approved. Between each pair of neighbouring shots is a JOIN: '
    'a Cut generates nothing, every other join (Continuous move, Dissolve, Match cut, Whip pan, Morph) generates one clip '
    'from the first shot\'s frame to the next. A shot belongs to ONE sequence only.',
    [('Name: ', 'click the title to rename it.'),
     ('The whole sequence: ', 'a description that reaches every clip.'),
     ('Shots and joins: ', 'each shot with × to take it out; between shots, the join type and "How 1B becomes 1C". '
      '"+ Add a shot…" lists every shot not yet in a sequence.'),
     ('Start and end frames: ', '"Starts on" / "Ends on" — its own first / last shot, or a frame borrowed from another sequence.'),
     ('How to make it: ', 'Leg by leg, One native clip, or Upload. "Join legs into one clip · free" stitches legs.'),
     ('Video versions: ', 'clips made for the whole sequence — the selected one plays once across all its shots.'),
     ('Footer: ', '"Plan · free" lists every clip that would be made; "Generate N clips" spends.')])
pics_side(['23-drawer-seq-a', '24-drawer-seq-b'], 'The Sequence drawer.', each=Inches(3.0))

doc.add_heading('5.3 Linked frame', level=2)
pic('12-node-link', 'A Linked frame: 1A\'s frame lent to "We\'re closed" as its start frame.', width=Inches(2.6))
node_card(
    'A frame borrowed from one sequence into another, so the second sequence starts (or ends) exactly where the first '
    'left off. It sits in the receiving sequence\'s box with a dashed border and is labelled with its source.',
    'It is never membership: the borrowed frame is used as a keyframe, and the source shot is not played twice. It always '
    'follows the SOURCE\'s selected version; when the source changes, the link reads "redo" until the receiving sequence '
    'is generated again.',
    [('Drawer: ', 'read-only — the picture, where it comes from and where it goes. "Unlink" removes it; '
      '"Open the source" jumps to the sequence it borrows from.'),
     ('Making one: ', 'drag "Linked frame" from the palette, or wire one sequence\'s "end frame" output into another '
      'sequence\'s "start frame" or "end frame" slot.')])
pics_side(['31-link-modal', '28-drawer-link'], 'Left: the "Link a frame" dialog. Right: the Linked frame drawer.', each=Inches(3.0))

doc.add_heading('5.4 Video version', level=2)
pic('13-node-video', 'A Video version: "v1 of 1", its source, and a thumbnail you can play.', width=Inches(2.6))
node_card(
    'One clip made for a shot or a sequence — generated, uploaded, or a stand-in. "▶ playing" means it is the selected one.',
    'Every generation adds a new version; nothing is overwritten. Exactly one version per parent is selected, and the '
    'selected one is what Playback, the NLE export and the final film use. A new clip becomes the selected one automatically.',
    [('Drawer: ', 'its state, the clip player, parent, generator and date, "How was this made" with "Make another like this", '
      'every sibling version with Open / Select, and "Compare with selected".'),
     ('Footer: ', 'Download, and "Select — play this version" (or Unselect).')])
pic('27-drawer-version', 'The Video version drawer.', width=Inches(3.0))

doc.add_heading('5.5 Sound (SFX, Room tone, Music)', level=2)
pics_side(['14-node-sound', '15-node-audio'], 'Left: a Sound node (a Music cue). Right: one of its audio versions, with the waveform.', each=Inches(2.6))
node_card(
    'A written sound for the film: a sound effect, a room tone / ambient bed, or music. The text is your direction. '
    'With no version yet, a dashed "No version yet" node sits to its right with Generate and Upload.',
    'A sound scores its whole scene by default. Wire a shot or a sequence\'s "scene" output into it and it scores only '
    'that span. Its files appear as audio versions to its right; the selected one plays in Playback and goes into the '
    'final film at its offset and level.',
    [('Kind: ', 'Sound effect, Room tone or Music.'),
     ('Written from / Your direction: ', 'the scene details it is written from, and what you add on top.'),
     ('Mood, Genre, Length, Avoid: ', 'music mood and genre; "Match" or "Custom" length; what it must not contain.'),
     ('What will be sent: ', 'the exact text the generator receives, and which provider.'),
     ('Versions: ', 'each audio version with a player, Open, Select / Unselect. "Open in Score workspace →" for the full score editor.'),
     ('Footer: ', '"Preview · free", Upload, Record, and Generate.')])
pics_side(['25-drawer-sound-a', '26-drawer-sound-b'], 'The Sound drawer.', each=Inches(3.0))

doc.add_heading('5.6 Audio version', level=2)
p('The audio counterpart of a video version: one generated or uploaded file for a sound, with its waveform, its length '
  'and "▶ playing" when selected. Its drawer is the same as a video version\'s, with an audio player.')

# ── 6 ─────────────────────────────────────────────────────────────────────
doc.add_heading('6. Live progress on the node', level=1)
p('While something is being generated, its node says so: a bar across the picture, what the provider is doing, the '
  'percentage when there is one, and how long it has been running. It is the same wherever the generation was started — '
  'from this page, from a batch run, or by Claude through MCP — because progress is written to the project, not held in '
  'the page.')
pic('42-node-progress', '1C while its clip is being made: "rendering · 62%" and the elapsed time, with a bar across the frame.', width=Inches(2.6))
table(['What the provider reports', 'What you see', 'Providers'], [
    ['A percentage', 'A filling bar, "62%", the phase and the elapsed time.', 'Runway, BFL, Meshy, Gridlight'],
    ['Only a phase', 'The phase ("queued", "processing") and the elapsed time; no bar.', 'MuAPI, Seedance, World Labs'],
    ['Nothing', 'The elapsed time and "no percentage from this provider".', 'the rest (they answer in one go)'],
], widths=[Inches(1.8), Inches(2.9), Inches(2.0)])
p('A percentage is never invented: if the provider has not said, the node says it has not said.')

# ── 7 ─────────────────────────────────────────────────────────────────────
doc.add_heading('7. Out-of-date on the graph, and "Run what changed"', level=1)
p('"Run what changed" (header) redoes only what is behind, in the right order, after one confirmation. It reads the '
  'same states the nodes show: every "redo" item, in the order the film is built (frames before the clips made from '
  'them), each with its price and a total.')
pic('49-run-changed', 'The confirmation: two frames to redo at $0.04 each, and "Left out" naming the held shot and why.', width=Inches(5.5))
bullets([
    ('What it leaves out, and says so: ', 'anything "waiting" (it runs once what it waits on is done), a card only a person '
     'can rewrite, frames on a locked board, and held nodes.'),
    ('One at a time: ', 'each item runs through its normal Generate, then the states are re-read before the next — so a '
     'sequence waiting on a frame runs as soon as that frame is redone.'),
    ('It stops at the first refusal ', '(a provider saying no, the budget) and names what it did not attempt. It never buys '
     'the same thing twice.'),
    ('Budget: ', 'if the total would take the project past its budget, the confirmation says so and nothing runs.'),
])

# ── 8 ─────────────────────────────────────────────────────────────────────
doc.add_heading('8. Run to here', level=1)
p('Right-click a shot, a video version\'s parent, a sequence or a sound and choose "Run to here": it makes everything '
  'that node still needs, in order, after one confirmation — the frames first, then the clip; for a sequence, every '
  'member frame, any borrowed frame traced back to its source, then the sequence\'s clips (a cut costs nothing).')
pics_side(['48-menu', '50-run-to-here'], 'Left: "Run to here" on the We\'re closed sequence. Right: its plan — redo 1C\'s frame, '
          'then make the sequence clip — and 1B named as held and left out.', each=Inches(3.0))
bullets([
    ('Blockers stop it before anything runs: ', 'a card only a person can rewrite, frames on a locked board.'),
    ('It stops at the first refusal ', 'and names what it did not attempt; progress shows on each node as it goes.'),
    ('Held nodes ', 'are listed as "Held, left out" — the plan runs around them, and says so.'),
])

# ── 9 ─────────────────────────────────────────────────────────────────────
doc.add_heading('9. Queue and history strip — and what Cancel really does', level=1)
p('Under the canvas, the queue strip lists every generation and flow run for this film in six groups. Closed, it shows the counts; '
  'click "Queue" to open it. Click any item to jump to its node.')
pic('47-queue', 'The strip open: one running (with Cancel), two done today, one waiting to be collected (Collect) and one failed (Re-run).')
table(['Group', 'What is in it', 'What you can do'], [
    ['Running', 'Being made now, with its percentage or phase.', 'Cancel — or Stop waiting (below).'],
    ['Waiting', 'The rest of a "Run what changed" or "Run to here" still to come.', 'Cancel run: stops before the next step.'],
    ['Waiting for a pick', 'A flow run that stopped at a "pick one" step. Its variations are on the shot, not yet on the board.', 'Pick — opens the shot, where each variation has a Pick button; Cancel run keeps the variations and ends the run.'],
    ['Done today', 'Finished today.', 'Click to see it on its node.'],
    ['Waiting to be collected', 'A job the provider accepted and nobody was waiting for any more (for example a call Claude gave up on after 60 seconds). The provider probably finished it and billed for it.', 'Collect — free; fetches the result onto its node.'],
    ['Failed', 'What went wrong, in the provider\'s words.', 'Re-run — through the node\'s own confirmation.'],
], widths=[Inches(1.4), Inches(3.3), Inches(2.0)])
doc.add_heading('What Cancel really does', level=3)
bullets([
    ('Runway ', 'really cancels at the provider: the job stops and is marked cancelled.'),
    ('Everyone else ', '(BFL, Meshy, MuAPI, Seedance, World Labs) cannot be stopped once they have the job. The button says '
     '"Stop waiting": Film Engine stops waiting, the job moves to "Waiting to be collected", and you are told the provider '
     'may still finish it and bill for it. You can still Collect it later.'),
    ('A run ', '(Run what changed, Run to here) is cancelled between steps: the step in progress finishes, nothing after it starts.'),
])

# ── 10 ────────────────────────────────────────────────────────────────────
doc.add_heading('10. How was this made, Make another, Compare', level=1)
p('Every version — a frame, a clip, a sound — has a "How was this made" panel in its drawer: the generator (provider and '
  'model), size, whether its inputs have changed since, the render-ledger entry, the cost, when it was made, and the '
  'prompt, negative prompt, references and seed when they were recorded. Anything not recorded is named ("Not recorded: '
  'prompt, seed…") rather than guessed or borrowed from another version. A cost or ledger entry found by time rather than '
  'by file says "matched by time".')
pic('46-drawer-howmade', '"How was v1 made" for 1A\'s frame, and "Make another like this".', width=Inches(3.4))
bullets([
    ('Make another like this: ', 'opens the usual confirmation already filled with that version\'s prompt — and its seed, '
     'where the provider honours one (the panel says when a seed is ignored). It makes a new version; nothing is replaced.'),
    ('Compare with selected: ', 'an A/B wipe between a version and the selected one — drag the slider across. For frames '
     'and clips.'),
    ('Drop a file to find its recipe: ', 'drag a picture, clip or sound from your computer onto the canvas. If it is one of '
     'this film\'s files (matched by its contents, which never leave your machine), its node is selected and its recipe '
     'opened; if not, it is offered as an upload to the node you dropped it on.'),
])
pic('55-compare', 'Compare with selected: v1 on the left, the selected v11 on the right, the slider in between (from Wingfall, shot 2A).', width=Inches(4.5))

# ── 11 ────────────────────────────────────────────────────────────────────
doc.add_heading('11. Hold a node', level=1)
p('Hold a shot, sequence or sound to keep it out of every batch run while you decide about it. Right-click → Hold, or '
  'select it and press Ctrl+B (Ctrl+B again, or "Release hold", releases it).')
pic('43-node-held', '1B held: the HELD tag and the dashed border.', width=Inches(2.6))
bullets([
    ('What skips it: ', 'Run pending, Run what changed, Run to here, and every batch generate (all frames, all videos, all '
     'dialogue, all music, lip-sync, post). Each one names what it left out and why.'),
    ('What does NOT skip it: ', 'the film. A held shot still plays in Playback, is still in the NLE export and the final '
     'master. Hold stops spending on it, never removes it.'),
    ('Generating it on purpose ', 'still works: its own Generate button is not a batch.'),
])

# ── 12 ────────────────────────────────────────────────────────────────────
doc.add_heading('12. Collapse a group', level=1)
p('Click the arrow at the top right of a sequence\'s or scene\'s box to collapse it into one card: its length, shots done '
  'out of the total, how many of its nodes are behind, and anything running. Click the arrow again to expand it.')
pics_side(['53-collapsed', '54-side-collapsed'], 'Left: "The empty diner" collapsed — 5 s, 1 of 1 shots done, 2 behind. '
          'Right: the side panel names the card a hidden shot is in.', each=Inches(3.0))
bullets([
    ('Nothing moves. ', 'Collapsing hides the nodes; expanding puts every one back exactly where it was.'),
    ('Tidy keeps it. ', 'Tidy layout never forgets the positions inside a collapsed group.'),
    ('Still findable. ', 'Fit and the minimap include the card; clicking a hidden shot in the side panel jumps to its card.'),
    ('Remembered ', 'per film, until you expand it.'),
])

# ── 13 ────────────────────────────────────────────────────────────────────
doc.add_heading('13. Search to add, and patterns (double-click)', level=1)
p('Double-click any empty part of the canvas to add by name, from the keyboard, where you clicked. Type to filter, ↑↓ to '
  'choose, Enter to add, Esc to cancel.')
pic('51-palette', 'The palette: a shot after the nearest shot (with its code), three coverage patterns, and every kind of cue.', width=Inches(4.0))
bullets([
    ('A shot: ', '"Shot after 1A — becomes 1AA": inserted after the nearest shot with the next insert code; no other shot is '
     'renamed. Enter, then type what happens in it.'),
    ('A cue: ', 'score, source music, sound effect, room tone or transition, for the scene you are in.'),
    ('A sequence: ', 'Shift+click shots on the canvas first, then double-click: "Sequence of 1B, 1C".'),
    ('A pattern: ', 'shot / reverse shot, insert then reaction, or wide, medium, close. Enter once shows a FREE preview of '
     'the shots and the sequence it would make; Enter again creates them. Patterns create shots and a sequence with its '
     'joins and generate nothing.'),
])
pic('52-pattern-preview', 'Shot / reverse shot previewed after 1C: 1CA over-the-shoulder on the speaker, 1CB the reverse, joined by a cut.', width=Inches(4.0))

# ── 14 ────────────────────────────────────────────────────────────────────
doc.add_heading('14. Wiring: what connects to what', level=1)
p('Drag from an OUTPUT port (right side) and drop on an INPUT port (left side). While you drag, the ports that will '
  'accept the wire light up. Only these connections exist:')
table(['Drag from', 'Drop on', 'What it does'], [
    ['A shot\'s "image" (purple)', 'A sequence\'s shot slot or "+ shot"', 'Adds the shot to the sequence. If it is already in another sequence you are asked whether to move it.'],
    ['A shot or sequence\'s "scene"', 'A sound\'s "scene" input', 'The sound now scores just that shot or sequence, and is written from it.'],
    ['A sequence\'s "end frame"', 'Another sequence\'s "start frame" or "end frame"', 'Makes a Linked frame.'],
], widths=[Inches(1.9), Inches(1.9), Inches(2.9)])
p('Versions are never wired by hand — the engine creates them and their wires when something is generated or uploaded.')

# ── 15 ────────────────────────────────────────────────────────────────────
doc.add_heading('15. Adding, duplicating and removing', level=1)
doc.add_heading('Adding', level=3)
p('Double-click the canvas (section 13), or drag an item from "Add node · drag onto graph" in the side panel.')
table(['Palette item', 'What it creates'], [
    ['Shot', 'A new shot after the last shot of the scene in view — you type what happens in it.'],
    ['Sequence', 'An empty sequence you name. Free.'],
    ['SFX / Ambient / Music', 'A new sound of that kind, scoring the scene in view.'],
    ['Linked frame', 'Opens "Link a frame" (needs at least two sequences).'],
], widths=[Inches(1.6), Inches(5.1)])
doc.add_heading('Right-click menu', level=3)
pic('30-menu', 'Right-click on a held shot.', width=Inches(1.8))
table(['Item', 'What it does'], [
    ['Generate frame / Regenerate frame, Generate sequence, Generate sound', 'The node\'s own generate — the same confirmation as its drawer button.'],
    ['Play this version', 'On a video or audio version: selects it.'],
    ['Open the source sequence', 'On a linked frame: jumps to the sequence it borrows from.'],
    ['Run to here', 'Everything this node still needs, in order, after one confirmation (section 8).'],
    ['Duplicate', 'Shot: a new shot after it with the same description. Sequence: a copy WITHOUT its shots. Sound: a copy with the same direction.'],
    ['Upscale…', 'On a shot, or one of a shot\'s clip versions: enlarges the clip to the delivery size on MuAPI, after one confirmation (section 27).'],
    ['Hold / Release hold', 'Batch runs skip it, or stop skipping it (section 11). Also Ctrl+B.'],
    ['Remove from its sequence', 'On a shot in a sequence: takes it out. The shot and its frames are kept.'],
    ['Remove from graph', 'Sequence: deletes the sequence — its shots and clips stay. Sound: deletes the cue — its audio stays. Linked frame: unlinks.'],
    ['Right-click a box: Run pending, Tidy layout', 'Run pending for what is in view; re-place the nodes nobody moved.'],
], widths=[Inches(2.3), Inches(4.4)])

# ── 16 ────────────────────────────────────────────────────────────────────
doc.add_heading('16. The header', level=1)
bullets([
    ('Counts: ', 'shots, how many are framed, how many have video, how many still need a frame.'),
    ('Legend: ', 'how many nodes are current, redo, waiting, not made and untracked; "Only what\'s behind (N)" filters the canvas.'),
    ('Scene chips: ', '"All scenes" or one scene.'),
    ('Tidy layout: ', 're-places every node you have not moved.'),
    ('Lock board: ', 'freezes the storyboard. While locked, nothing that would replace a frame runs.'),
    ('Run what changed: ', 'redo only what is behind, in order (section 7).'),
    ('Run pending (N): ', 'every shot without a frame and every sound without a version, in view, after ONE confirmation. '
     'Held nodes are named and skipped. "Nothing pending" when everything has been made.'),
])

# ── 17 ────────────────────────────────────────────────────────────────────
doc.add_heading('17. The playback bar', level=1)
pic('04-playback', 'The playback bar: monitor, transport, timecode, a tile per shot, and a lane per sound.')
bullets([
    ('Monitor: ', 'the current shot — its selected video, or its selected still held for the shot\'s length, or a slate.'),
    ('Transport: ', 'previous shot, play / pause, next shot, and the timecode.'),
    ('Tiles: ', 'one per shot in running order, sized by length. Click a tile to jump there and select that shot.'),
    ('Sound lanes: ', 'MUSIC, SFX and AMB, one lane per sound.'),
    ('Stills hold for the shot\'s length / Audio: ', 'switches for how the preview plays.'),
    ('Send to Premiere: ', 'opens the Export page (section 25).'),
])

# ── 18 ────────────────────────────────────────────────────────────────────
doc.add_heading('18. What costs money, and what is free', level=1)
p('Everything that spends goes through one confirmation that shows the provider, model, what will be sent and the '
  'estimated cost before anything is bought. Every plan is free.')
table(['Free (nothing is generated)', 'Spends credits'], [
    ['Moving, wiring, selecting versions, Tidy, Lock, Hold, Collapse', 'Generate / Regenerate / Refine a frame'],
    ['Creating a sequence, a sound, a link, or a pattern\'s shots', 'Generate video (a shot)'],
    ['Preview request, Plan, Preview (sound), a pattern\'s preview', 'Generate a sequence (per clip that is not a cut)'],
    ['The plans of Run what changed and Run to here', 'Generate a sound, Generate dialogue'],
    ['How was this made, Compare, dropping a file to find its recipe', 'Run pending, Run what changed, Run to here (what they list)'],
    ['Collect (the provider already billed for it)', 'Make another like this, Re-run'],
    ['Uploading your own frame, clip or audio', 'Upscale… (MuAPI)'],
    ['The delivery check, Premiere by scene, Back up now', ''],
], widths=[Inches(3.35), Inches(3.35)])
p('Stop waiting does not un-bill a job: the provider may still finish it and charge for it, which is why the button says so.')

# ── 19 ────────────────────────────────────────────────────────────────────
doc.add_heading('19. Asking Claude to do it', level=1)
p('Everything on this page can be done by Claude through Film Engine\'s MCP connection — no separate API key; your Claude '
  'plan does the thinking. The plans are free; the runs spend, exactly as they do here. Ask in plain words ("what is behind '
  'on Last Call?", "hold 1B", "run what changed"); these are the tools Claude uses:')
table(['Tool', 'What it does'], [
    ['production_graph_get', 'The whole graph: nodes, states, holds, groups, what is running.'],
    ['production_graph_running', 'What is running right now, with progress.'],
    ['generation_queue', 'The queue strip: running, waiting, done today, waiting to be collected, failed.'],
    ['run_changed_plan · run_changed · run_changed_status', 'Plan (free), run (spends) and follow "Run what changed".'],
    ['run_to_here_plan · run_to_here', 'Plan (free) and run (spends) "Run to here" for a node.'],
    ['run_cancel · generation_cancel', 'Cancel a run between steps; cancel or stop waiting for a job.'],
    ['asset_provenance', 'How a version was made.'],
    ['pattern_list · pattern_preview · pattern_create', 'The coverage patterns, their free preview, and creating them.'],
    ['graph_hold', 'Hold or release a shot, sequence or sound.'],
    ['version_select', 'Choose which version plays.'],
    ['flow_apply_plan · flow_form', 'Plan applying a flow to shots, and read its form — both free.'],
    ['flow_apply · flow_apply_get', 'Apply the flow with the plan\'s fingerprint (spends), and follow the apply.'],
    ['flow_run_select · flow_run_cancel', 'Pick a paused run\'s variation, or cancel a run.'],
], widths=[Inches(2.6), Inches(4.1)])

# ── 20 ────────────────────────────────────────────────────────────────────
doc.add_heading('20. Flows on the canvas', level=1)
p('A flow is a small graph of generations built on the Flows page: a prompt, a keyframe, three video models, "pick a take". '
  'On the canvas you can run one on as many shots as you like, see what it made on each shot, and choose what stays.')
doc.add_heading('Apply a flow', level=3)
bullets([
    ('Apply flow… ', 'is on every node\'s right-click menu. It applies to the shots you picked (Shift+click), or to the one you selected; '
     'a sequence means its shots. On a clip or a sound it is greyed out and says where to apply instead.'),
    ('The picker ', 'lists this film\'s flows, your library flows and the six ready-made templates. Choosing a template saves it into the film first.'),
    ('The form. ', 'If the flow has inputs marked "exposed", you fill them in here (a prompt, a picture, a character) for this apply only.'),
    ('The plan ', 'is free and comes before anything spends: what each shot will use, what it costs, the total against the budget, '
     'and any shot that is held or cannot take the flow. Generate stays greyed out until the plan has loaded.'),
    ('Templates on the shelf. ', 'Double-click the canvas: the add palette lists the templates beside the patterns, and choosing one applies it.'),
])
doc.add_heading('What a flow makes', level=3)
bullets([
    ('Made by flows. ', 'A picture a flow made sits on its shot in a "Made by flows" strip; a clip or a sound becomes a version with a flow badge. '
     'None of it replaces anything: the board, Playback and the final film keep what they had.'),
    ('Pick. ', 'When a flow stops at "pick one", every variation waits on the shot with a Pick button. Picking makes it the shot\'s frame, '
     'clip or sound and finishes the run from there; nothing is generated again. Cancel run keeps the variations and ends the run.'),
    ('Waiting for a pick ', 'is its own group in the queue strip; each entry opens its shot.'),
    ('How was this made ', 'names the flow and its version, the node that made it, the provider and model, the run and the apply. '
     'Make another applies the same flow to the shot again.'),
])
doc.add_heading('Editing the flow', level=3)
bullets([
    ('Edit flow ', 'opens the Flows page on that flow, from the picker, the "Made by flows" strip, a flow version, the queue or the form. '
     'Back (the browser\'s, or the page\'s own) returns to the canvas.'),
    ('Suggestions. ', 'On the Flows page, drag a wire from an output and let go on empty space: it lists only the nodes that can take it, '
     'and choosing one places it there, wired.'),
])

# ── 21 ────────────────────────────────────────────────────────────────────
doc.add_heading('21. The Score page', level=1)
p('A Music node\'s "Open in Score workspace →" opens the Score page (Production → Score). A score session is written against '
  'a picture: a sequence, a scene, or an edit brought back from Premiere (section 23). Lanes, clips, the shots, hit markers '
  'and the emotional arc sit on one ruler with one playhead. Playing it in the page is free: nothing is generated to hear it.')
doc.add_heading('Lanes and their details', level=3)
bullets([
    ('Clips: ', 'drag to move, drag an edge to trim (the file is never cut). Takes in a group are one click apart.'),
    ('The chevron beside a lane\'s name ', 'opens its details underneath it, in six tabs: Part, Instrument, Mix, Takes, '
     'Automation, Lane. The side panel only says which lane is selected.'),
    ('Part is a piano roll ', 'lined up with the clips above it. Click to add a note, drag to move, drag its right edge for '
     'length, double-click to remove. Snap is 1/4, 1/8, 1/16 or off, from the session\'s own tempo. Clicking a key plays a '
     'plain reference tone, not your instrument.'),
    ('Play ', '(on the lane head) plays the lane\'s notes through its instrument (section 22). It is free, and the result '
     'lands as a new take that is auditioned at once; what was selected stays selected until you choose otherwise. '
     'When a lane cannot play (no notes, no instrument, a plug-in that has gone), the button is greyed out and says why.'),
    ('Import stems ', 'lays a composer\'s stems on new lanes, all starting at the playhead. The originals are stored '
     'untouched, with an optional 48 kHz working copy, and their rights recorded as declared (or "unknown").'),
])
doc.add_heading('AI actions and Jobs', level=3)
bullets([
    ('AI actions: ', 'Emotional arc: proposals, Whole score, Score the selection, Native parts, Separate into stems, '
     'From this as a reference, From the picture, Inpaint at the playhead, Regenerate as a new take, A/B audition, '
     'Approve this take. Anything that spends shows its free plan first in the usual confirmation. An action the '
     'project\'s music provider cannot do is greyed out with that provider\'s own reason.'),
    ('The emotional arc ', 'is proposed by Claude and accepted by you, range by range. Nothing is generated from a proposal '
     'nobody accepted.'),
    ('Every output is a new take. ', 'Nothing replaces what you had; a take becomes the one that plays only when you approve it.'),
    ('Jobs: ', 'each generation or separation with its parts. "Check again" on one still running, Retry (spends, through the same confirmation) on one that failed.'),
])
doc.add_heading('Approving the score', level=3)
bullets([
    ('A bounce ', 'is the session rendered to a master and stems — ask Claude ("bounce the score"); an unchanged session is '
     'not rendered twice.'),
    ('Approve newest bounce ', 'makes it the film\'s score. A bounce older than your last change is refused. From then on '
     'the mix plays once at its place in Playback, the Premiere export and the final film, and the scene music under it is '
     'dropped. Rights marked blocked stop the approval.'),
    ('Ableton Live (optional): ', 'the DAW panel pushes lanes to Live and reads the set, through a small helper you start '
     'yourself (docs/ableton-sidecar.md). Every DAW action has a twin that needs no DAW: Export a score package, Import a '
     'package, Play in the workstation.'),
])

# ── 22 ────────────────────────────────────────────────────────────────────
doc.add_heading('22. Your own instruments', level=1)
p('Your sample libraries (Kontakt and other plug-ins) can play a lane\'s notes inside Film Engine, with no DAW open and '
  'nothing billed. An instrument here is one sound: the plug-in plus the saved patch that recalls it, named from Kontakt\'s '
  'own index ("Vortex Bells, Ethereal Earth") so you can always tell where a file came from.')
bullets([
    ('The instrument helper ', 'is a small program you start by hand on the Mac (docs/instrument-sidecar.md). It only loads '
     'plug-ins from the folders macOS installs them in.'),
    ('The library holds only what you used. ', 'Browsing your sounds reads Kontakt\'s index live; a sound joins the library the '
     'first time it plays a part.'),
    ('It renders, it does not play live. ', 'A part is rendered far faster than real time and comes back as a take you hear '
     'on the Score page.'),
    ('A silent render is refused ', 'rather than kept — silence is what a plug-in with no patch loaded produces.'),
    ('Render with instruments ', '(the MIDI panel on a cue in Music & Sound) plays a cue\'s notes through a SoundFont on this '
     'Mac instead; its licence is recorded on the file.'),
])

# ── 23 ────────────────────────────────────────────────────────────────────
doc.add_heading('23. An edit made in Premiere', level=1)
p('The Edit page (Post) brings a cut finished in Premiere back into Film Engine, so a score can be written against the film '
  'that will actually play rather than the assembly.')
bullets([
    ('Import an edit… ', 'takes the picture you exported from Premiere (H.264 or ProRes). Each import is a new version '
     '(edit_v1, edit_v2…) in the project\'s 05 Edit folder; nothing is overwritten.'),
    ('Then the cut list: ', 'Final Cut Pro XML (File → Export → Final Cut Pro XML) or an EDL. Each event is matched to its shot '
     'by the clip file, else by shot code; a title or stock shot matches nothing and is kept, named.'),
    ('Score this edit ', 'opens a score session exactly the edit\'s length, so its stems drop onto the Premiere sequence at 00:00.'),
    ('A newer edit ', 'is reported on the session and applied only when you ask.'),
])

# ── 24 ────────────────────────────────────────────────────────────────────
doc.add_heading('24. The final film and its sound', level=1)
p('Conform joins each shot\'s selected clip in running order into one master file, on this Mac, for free. It needs every '
  'shot to have a clip: a missing one is named and nothing is built.')
table(['Sound', 'Where it goes in the master'], [
    ['Dialogue', 'Each shot\'s lines from the shot\'s start, with the card\'s pauses — only under a clip that has no sound of its own.'],
    ['Scene music and room tone', 'At the cue\'s offset in its scene, with its level and fades, stopped where the scene ends.'],
    ['Sound effects', 'From the start of the shot they belong to, at −4 dB.'],
    ['The approved score', 'Once, at the start of its picture, in place of the scene music it covers.'],
    ['A finished project mix', 'If one is registered, it is the soundtrack and nothing is laid over it.'],
], widths=[Inches(1.8), Inches(4.9)])

# ── 25 ────────────────────────────────────────────────────────────────────
doc.add_heading('25. Sending the film to Premiere', level=1)
bullets([
    ('Send to Premiere ', '(Export page) downloads a Premiere Pro XML: one sequence in running order, each clip named by its shot '
     'code with a "Scene N" marker, and separate audio tracks for dialogue, music, sound effects and ambience. It points at '
     'the files where they are on this Mac. In Premiere: File → Import.'),
    ('Premiere, by scene ', '(Export page, "Make the folders") is the handover an editor opens: in 07 Delivery/Exports, a folder '
     'per scene (Scene_01_INT-DINER-NIGHT…) with each shot\'s selected clip in Video, named by its shot code, and its dialogue, '
     'effects and the scene\'s beds in Sound. The XML beside them carries the cut and a bin per scene, and points at the copies, '
     'so everything opens online. A READ ME lists anything missing. Free; the originals are copied, never moved.'),
    ('Every export uses the clip you selected. ', 'It used to take a shot\'s oldest clip; now it is the same one Playback and the '
     'final film play.'),
    ('Shots with no clip ', 'are left out of the Premiere sequence; the export preflight names them, and warns about any clip '
     'below the delivery size (section 26).'),
    ('A packaged handover ', '(ask Claude for export_package) writes the XML with a copy of every file it uses into one folder '
     'under 07 Delivery/Exports, for another machine or an editor.'),
])

# ── 26 ────────────────────────────────────────────────────────────────────
doc.add_heading('26. Delivery quality: the resolution you asked for', level=1)
p('Settings → Technical Settings decides what the film is delivered at. Every generator is asked for that size; one that cannot '
  'reach it renders the best it can, and says so before you pay.')
bullets([
    ('Resolution: ', 'every clip is asked for it. Runway now renders 1080p on a model that offers it (it used to send 720p every '
     'time). A generator whose best is smaller shows "below delivery — its best; upscale after" in the video confirmation.'),
    ('Draft video ', '(a checkbox, off by default): clips at the model\'s cheapest size while you are blocking, upscaled at the end. '
     'It used to be on for every project with no way to see it.'),
    ('Delivery preset: ', 'choosing one now saves its codec and audio layout too. Delivery codec (H.264, H.265, ProRes 422 Proxy '
     'to 4444, DNxHR HQ, JPEG 2000) and Delivery audio (stereo, 5.1, 7.1, 12.0) can also be set by hand.'),
    ('The final film ', 'is always an H.264 stereo master; when the delivery asks for anything else, Conform also writes a '
     'delivery master beside it in that codec and layout.'),
    ('The delivery check ', '(top of the Export page, free) measures every shot\'s selected clip from the file and lists any '
     'below the delivery size, with the fix: Upscale (section 27).'),
])

# ── 27 ────────────────────────────────────────────────────────────────────
doc.add_heading('27. Upscaling a clip', level=1)
p('Right-click a shot or one of its clip versions on the canvas and choose "Upscale…" (on the Video Shots page it is the '
  'Upscale button). The confirmation measures the clip, names the size it will reach, and lets you pick the upscaler, '
  'each priced for this clip. They come from three places:')
table(['Upscaler', 'Where', 'What it does'], [
    ['Topaz video upscale (default)', 'MuAPI', '2x or 4x; about $0.08 a second'],
    ['AI video upscaler / Pro', 'MuAPI', 'to 1080p, 2K or 4K'],
    ['FLUX.3 video upscaler', 'MuAPI', 'prompted, 2x to 4x'],
    ['Starlight Precise 2.6', 'Topaz', 'the best finish for generated footage, to 4K; about $1.44 for 10s at 1080p'],
    ['Starlight Fast 3', 'Topaz', 'the same quality, up to four times faster, same price'],
    ['Astra 2', 'Topaz', 'creative: adds new detail, can be steered with a prompt'],
    ['Proteus', 'Topaz', 'precision: keeps the picture as it is; cheapest'],
    ['Magnific Video Upscaler / Turbo', 'Magnific', 'creative: creativity and flavor (vivid or natural)'],
    ['Magnific Video Upscaler Precision', 'Magnific', 'faithful: a strength that blends original and upscale'],
    ['Topaz Starlight via Magnific', 'Magnific', 'Starlight on a Magnific account, for those without a Topaz key'],
], widths=[Inches(2.2), Inches(0.9), Inches(3.6)])
bullets([
    ('Topaz and Magnific need their own API key ', '(Settings, AI providers and models). Until one is set their upscalers are '
     'listed, greyed out, with the reason.'),
    ('Prices: ', 'MuAPI lists its prices without a unit, so they are counted per second of the clip, which errs high. Topaz '
     'prices come from its own credit table ($0.12 a credit on the Starter plan) at the size the clip reaches. Magnific '
     'publishes no per-frame rate, so its prices are the ones Runway resells it at.'),
    ('The result ', 'is a new version that keeps its sound and becomes the clip that plays. The original stays.'),
    ('The clip is uploaded first, ', 'for free, to whichever service does the work, and only then enlarged.'),
    ('Cancel: ', 'Topaz really stops (work not yet processed is refunded); MuAPI and Magnific cannot be stopped, so '
     'cancelling there means you stop waiting and the result can still be collected.'),
])

# ── 28 ────────────────────────────────────────────────────────────────────
doc.add_heading('28. Your recorded dialogue in a clip', level=1)
p('If you have recorded the lines yourself, Seedance 2.5 can take them as audio references so the performance follows your '
  'recording.')
bullets([
    ('Upload each line ', 'with "Upload dialogue": on each shot of the Video Shots page, and in the shot\'s drawer on the canvas. '
     'Name files so their order shows (1A_RAY_1.mp3, 1A_JUNE_2.mp3): they are sent in that order.'),
    ('Then, in the video confirmation, ', 'tick "Send this shot\'s recorded dialogue". Only Seedance 2.5 takes it: on MuAPI, or '
     'the seedance2_5 model on Runway. Any other model says it cannot, rather than ignoring the lines.'),
    ('On MuAPI ', 'a shot sent with its dialogue runs Seedance\'s reference workflow, where the storyboard frame guides the clip '
     'rather than being its exact first frame. The confirmation says so.'),
    ('Up to ten lines ', 'per clip. They are uploaded to MuAPI for free before the clip is asked for.'),
])

# ── 29 ────────────────────────────────────────────────────────────────────
doc.add_heading('29. Backups', level=1)
p('Settings → Backups. Choose a folder (Dropbox, Google Drive and a NAS all work) and Film Engine backs up its database '
  'there every 6 hours, or as often as you set. Each person who uses Film Engine sets their own folder; backups go in a '
  'sub-folder named for them and their Mac, so several people can share one folder.')
bullets([
    ('What is in a backup: ', 'the whole database, as a snapshot that is safe to take while Film Engine runs. Optionally each '
     'project\'s records as a separate file, so one film can be restored alone.'),
    ('What is not: ', 'the media. It lives in each project\'s own folder (Settings, "Where new projects are saved"), '
     'which can itself be in a shared or synced folder.'),
    ('Keep ', 'is how many backups are kept; older ones Film Engine wrote are deleted, and nothing else in the folder is touched.'),
    ('Back up now ', 'writes one immediately. To restore, stop Film Engine and copy a backup over the database; the steps are in '
     'latest.json in the backup folder, and Claude can read them to you (backup_folder_status).'),
])

# ── 30 ────────────────────────────────────────────────────────────────────
doc.add_heading('30. Setup: providers, models and keys', level=1)
p('Settings shows one section at a time, chosen from the list on its left. The first four belong to the open project: '
  'AI providers and models, Project, Picture and delivery, Where it is saved. The rest belong to this Mac: API keys, Backups, '
  'You and connection, Production graph. The section you last opened is the one you come back to. Longer explanations sit '
  'behind a small "Why" you can open.')
bullets([
    ('Image quality ', '(the row of buttons above the table): Draft, Standard, Precision or Auto. It chooses the image provider '
     'and model when neither is pinned below.'),
    ('Provider and model for each kind of thing, grouped as Pictures (image, 3D model, world), Motion (video, upscale and '
     'finish, lip-sync), Sound (voice, music, sound effects, ambience) and Writing: ', 'one card each. '
     'Pick the provider, then the model: the model menu lists only what that provider offers. Each card\'s last line says who '
     'chose it, and warns when nobody did and when the provider has no key. '
     '"Automatic" uses the provider already in use (named in the menu); "Default" lets the provider, or for images the quality '
     'tier, choose. Save providers and models to keep it.'),
    ('A model you pin is used for every generation of that kind ', 'in this project, unless you pick a different one in a '
     'generation\'s own confirmation. A model the provider does not offer is refused when you save, with the list it does offer.'),
    ('Text / LLM has no model menu: ', 'the connected agent (Claude or ChatGPT Desktop) is the model.'),
    ('API keys ', 'have their own section: one tile per provider saying what it is used for and whether its key is set, '
     'entered once for this Mac and shared by every project. "Get key" opens the provider\'s page.'),
])

# ── 31 ────────────────────────────────────────────────────────────────────
doc.add_heading('31. A video model\'s own options', level=1)
p('When you generate a clip, the confirmation shows the options of the model it will run on, for this clip only. '
  'Change the generator or the model and the options change with it; change an option and the price is worked out again. '
  'Anything left on Default is what the project would have used: the shot\'s own length, the frame closest to the project\'s.')
table(['Option', 'What it does', 'Where'], [
    ['Length', 'Seconds of footage. Seedance 2.5 runs 4–30s; Gen-4.5 2–10s; Veo 3.1 4, 6 or 8s', 'every model'],
    ['Frame', 'The exact size and shape, from the provider\'s own list (Runway Seedance 2.5 has 18)', 'every model'],
    ['Resolution', 'The quality tier (Hailuo 3: 768p or 2K). On MuAPI Seedance the model you pick IS the resolution', 'some models'],
    ['Generate sound', 'The model\'s own soundtrack. On by default at Runway; Film Engine scores the film separately', 'Seedance, Veo on Runway'],
    ['Ends on', 'Another shot\'s board frame, which the clip arrives at. Your board frame is always the first', 'Seedance, Veo'],
    ['Output format', 'MP4, ProRes, and the HDR formats: HDR10, HLG, HDR ProRes, EXR', 'Gen-4.5 on Runway'],
    ['Avoid', 'What should not appear', 'Veo 3.1'],
    ['Seed', 'Repeat a take as closely as the model allows', 'most models'],
], widths=[Inches(1.3), Inches(4.0), Inches(1.4)])
bullets([
    ('Every option comes from the provider itself: ', 'Runway\'s published specification and MuAPI\'s own answers, dated in the dialog. '
     'A value a model does not take is refused by name before anything is spent.'),
    ('HDR: ', 'MuAPI offers none for Seedance 2.5. On Runway it is an output format of Gen-4.5.'),
    ('Sound on MuAPI: ', 'MuAPI takes no sound switch for Seedance 2.5, so it decides whether the clip has Seedance\'s own sound.'),
])

# ── 32 ────────────────────────────────────────────────────────────────────
doc.add_heading('32. Changing a project\'s stage', level=1)
p('On the project list, click the stage badge on a project (Concept, Script, … Complete). It becomes a dropdown; choose the '
  'stage and it is saved straight away, without opening the project. Settings → Project Settings has the same choice.')

# ── 33 ────────────────────────────────────────────────────────────────────
doc.add_heading('33. Good to know', level=1)
bullets([
    ('Nothing is overwritten. ', 'Every generation is a new version; you choose which one plays.'),
    ('Red means do it now; amber means not yet. ', '"waiting" work is built on something that is itself being redone.'),
    ('Untracked is not the same as fine. ', 'It means Film Engine cannot tell — usually an upload.'),
    ('One shot, one sequence. ', 'Wiring a shot into a second sequence asks to move it.'),
    ('A cut generates nothing. ', 'Only joins that move make clips.'),
    ('Hold stops spending, not the film. ', 'A held shot is still in Playback, the export and the master.'),
    ('Cancel is honest. ', 'Runway and Topaz stop at the provider; elsewhere you stop waiting and the job can still be billed and collected.'),
    ('The final film uses what is selected. ', 'The conform joins each shot\'s selected clip in running order with the '
     'dialogue, score and ambience where Playback plays them.'),
    ('Previs decisions are read here, changed in Previs. ', 'The Shot drawer shows them; "Open in Previs" to change them (section 34).'),
    ('Lighting has three levels. ', 'The shot\'s own technique wins, then the location\'s; the mood board\'s style is the film\'s general look.'),
    ('The score is written against a picture. ', 'Approve it once and it plays everywhere the film does (section 21).'),
    ('Your instruments are free to play. ', 'A render through your own library costs nothing and lands as a take (section 22).'),
    ('Every generator is asked for your resolution. ', 'A smaller result is said before you pay and listed by the delivery check (section 26).'),
])

# ── 34 ────────────────────────────────────────────────────────────────────
doc.add_page_break()
doc.add_heading('34. Previs: how it works', level=1)
p('Previs (Plan → Previs) is where a shot is staged and lit in a 3D set before anything is paid for. You choose who is '
  'in the shot, where the camera stands, the lens, the move and the light, and see the result at once. Everything here '
  'is free. Nothing reaches the storyboard until you Apply it, so you can try as many angles as you like.')
pic('60-previs-console', 'The Previs console: shots on the left, the decisions strip, the view in the middle with the '
    'move timeline under it, and the Camera / Light / Direct / Explore / Scene tabs on the right.')

doc.add_heading('The set', level=3)
p('A shot is framed inside its location\'s world. The world is the set built for that location: built free in Blender '
  'from the location\'s plates, scanned with an iPhone, or imported as a GLB. It is in metres, so a 24mm lens and a '
  '1.55 m camera height mean what they say. "scale calibrated" beside the title confirms that. "Build set" (top right) '
  'makes or rebuilds it; see "Building the set" below.')

doc.add_heading('The four views', level=3)
pics_side(['62-previs-view-look', '63-previs-view-geometry'], 'Look: the set as the camera sees it, with real light. '
          'Geometry: the plain shapes the image model is given as the plate.', each=Inches(3.2))
pics_side(['64-previs-view-depth', '65-previs-view-plan'], 'Depth: near is light, far is dark. Plan: the set from '
          'above, with the camera, its field of view and everyone staged.', each=Inches(3.2))
bullets([('Look ', 'shows the set from exactly the shot\'s camera. 360° looks around from where the camera stands.'),
         ('Geometry and Depth ', 'show what generation receives: the shapes and distances, not the colours.'),
         ('Plan ', 'is the bird\'s-eye view. The blue wedge is what the lens sees. Scroll to zoom, drag to pan.'),
         ('Moving the camera with the mouse: ', 'in Look, Geometry and Depth, scroll to dolly forward and back (Shift: '
          'faster), drag to look around, and right-drag or Shift-drag to slide sideways and up or down.')])

doc.add_heading('Staging people, furniture and your own models', level=3)
pics_side(['66-previs-plan-staging', '67-previs-add-modal'], 'On the Plan, "+ Add people, furniture or a model" opens '
          'the library; the selected figure gets turn buttons, Frame on and Remove.', each=Inches(3.2))
p('People are a man, a woman, a boy and a girl at real heights; furniture is 140 library pieces; "Your 3D models" '
  'lists the project\'s own (a Meshy creature, say). Each is placed a few metres in front of the camera. Drag it on the '
  'Plan to move it, use the turn buttons to rotate it, "Frame on" to make it the subject the shot is framed on, and '
  'Remove to take it out. Staging never moves the camera.')

doc.add_heading('The decisions strip', level=3)
pic('61-previs-decisions', 'Each decision says whether it is only being tried here, applied to the shot, or not set.')
p('Camera, Direction, Lighting, Set view, Cast, Props and Move each carry a state: "trying" (staged here, not on the '
  'shot yet), "applied" (written to the shot\'s card), or "not set". "Apply N to the shot" writes what you are trying '
  'onto the card, which is what the storyboard frame is generated from. "Lock shot" freezes the applied decisions so '
  'later experiments cannot change them by accident. The same states appear in the Production canvas\'s Shot drawer '
  'under "From Previs" (section 5.1).')

doc.add_heading('Camera tab: framing in film terms', level=3)
pic('68-previs-camera-tab', 'Shot on, framing, angle, lens and Camera Operate.', width=Inches(2.6))
bullets([('Shot on ', 'chooses who the shot is framed on, from everyone staged. "Rename / move" names or repositions them.'),
         ('Framing ', 'EWS, WS, FS, MWS, MS, MCU, CU, ECU. The camera is moved so that subject fills the frame that way '
          'with the current lens, keeping the side you shoot from.'),
         ('Angle ', 'Eye level, Shoulder, Hip, Knee, Ground (worm\'s eye), Low, High, Overhead (bird\'s eye), Dutch.'),
         ('Lens ', '12 to 135 mm. "Keep position" holds the camera and the subject size follows the lens; "Maintain '
          'size" moves the camera so the subject stays the same size (a dolly-zoom).'),
         ('Camera Operate ', 'every number (distance, height, pan, tilt, roll, lens) can be typed, or dragged left and '
          'right with the mouse to change it while watching the view.')])

doc.add_heading('Light tab', level=3)
pics_side(['69-previs-light-tab', '70-previs-look-lit'], 'A technique, the side the key comes from and a mood, lit in '
          'the Look view.', each=Inches(3.2))
pic('71-previs-plan-lights', 'The Plan draws the lights: here the Rembrandt key and its fill around the subject.')
p('Pick a technique (Three-point, Rembrandt, Loop, Butterfly, Clamshell, Split, Broad, Short, Backlight / rim, '
  'Silhouette, Top light, Under light, Window, Practicals, High key, Low key, Chiaroscuro), the side the key light comes '
  'from, and a mood (golden hour, night, neon, candlelight…). The lights are placed from the saved camera on the subject, '
  'shown in Look ("Lights on in Look") and drawn on the Plan. Each line says where its value comes from: this stage, the '
  'shot, or the location. "Clear staged" goes back to what the shot or location says. Apply writes it to the shot, '
  'and the storyboard frame is described with it ("Rembrandt lighting, key light from camera left…").')

doc.add_heading('A camera move', level=3)
pics_side(['72-previs-timeline', '73-previs-look-mid-move'], 'The move timeline with two keys, and the view at the '
          'playhead part-way through the move.', each=Inches(3.2))
p('Put the camera where the move starts and press "+ Key"; move the playhead, put the camera where it ends and press '
  '"+ Key" again. Drag the playhead to see any moment of the move in the view; "Play move" plays it. "+ Leg" adds a '
  'named movement (dolly in, pan left…). Linear, Ease in, Ease out, Ease both and Hold shape how it speeds up and slows down.')

doc.add_heading('Walk', level=3)
pic('74-previs-walk', 'Walking: W/S/A/D to move, R/F for height, arrows or drag to turn, Shift for faster. '
    '"Use this camera" keeps it; "Put it back" returns to the saved camera.')
p('Walk lets you move through the set as if holding the camera. Nothing is saved while you walk. "Use this camera" keeps '
  'it as a checked proposal: a camera inside a wall is refused, exactly as any other change would be.')

doc.add_heading('More room', level=3)
pic('75-previs-folded', 'With both side panels folded away the view takes the whole page.')
p('The three buttons at the top right (◧ ◨ ⬓) hide or show the shot list, the right-hand tabs and the move timeline. '
  'This browser remembers the choice.')

doc.add_heading('Building the set', level=3)
pic('76-previs-build-set', '"Build set": ask Claude to build from the plates, or scan with an iPhone. Each attempt is '
    'kept, compared against the plates.', width=Inches(4.2))
bullets([('From the plates: ', 'copy the line under "Ask Claude" into Claude. It reads the plates, writes the room, and '
          'Blender builds and renders it beside each plate so you can see where it is wrong. Free, on this Mac.'),
         ('With an iPhone Pro (works offline): ', 'open the Film Engine app and tap "Scan a location". Scan a room '
          'while RoomPlan draws its lines live, tap "Room done", walk to the next room (or up the stairs) and "Scan another '
          'room". Then "Take photos": the app shows where to stand and which way to turn for each wall ("Walk 1.4 m ahead, '
          'turn left 20°"), and keeps each photo with the exact position of the phone. "Extra" takes any detail you want. '
          '"Done, save" keeps it on the phone. Back in reach of the Mac, tap "Send to Film Engine…" on the saved scan, '
          'choose the project and the location (or type a new one). The scan becomes the set, the photos become the '
          'location\'s plates, and the scan is rendered beside each photo.'),
         ('Then ask Claude ', 'to finish the set from its scan and photos: it keeps the measured walls and cameras and adds '
          'every object the photos show, as shapes and flat colours.'),
         ('Import RoomPlan JSON ', 'takes a scan made elsewhere.'),
         ('Finish an attempt ', 'to make it the location\'s world and a 3D model.')])

doc.add_heading('From Previs to the storyboard', level=3)
bullets(['Stage the cast, choose framing, angle and lens, light it, build the move.',
         'Press "Apply N to the shot". The card now says what you staged.',
         'Generate the frame from the Storyboard or the Production canvas: it uses the applied camera, cast and lighting. '
         '"Render plate" here saves the geometry and depth plate, from this camera, that the model is given.',
         'Lock the shot when it is right.'])
p('Claude can do all of this too: previs_get, previs_stage, previs_library, previs_timeline, previs_director (lighting and '
  'direction), previs_apply, previs_lock, set_build_brief, set_build_render and room_scan_import.')

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT, os.path.getsize(OUT))
