#!/usr/bin/env python3
"""
Render docs/using-film-engine.md to ~/Desktop/Film-Engine-Guide.docx.

The previous guide was a loose .docx on the Desktop with no source anywhere,
and it disappeared — with nothing to regenerate it from. The markdown is the
source now, versioned beside the code it describes; this only renders it.

    python3 scripts/export-guide.py

Needs python-docx. Nothing in the backend depends on this; it is a doc tool.
"""

import re
import os
import sys

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("python-docx is not installed:  pip3 install python-docx")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'docs', 'using-film-engine.md')
OUT = os.path.expanduser('~/Desktop/Film-Engine-Guide.docx')
ACCENT = RGBColor(0x2D, 0x4E, 0xA8)


def rich(par, text):
    """Bold, `code` and italics, without pulling in a markdown library."""
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8A, 0x30, 0x30)
        elif tok.startswith('*') and tok.endswith('*'):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def main():
    lines = open(SRC).read().split('\n')
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.85)
        s.left_margin = s.right_margin = Inches(0.95)

    i, first_h1 = 0, True
    while i < len(lines):
        ln = lines[i]

        if ln.strip() == '---':
            i += 1
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', ln)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            if lvl == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                r.bold, r.font.size, r.font.color.rgb = True, Pt(26), ACCENT
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sr = sub.add_run('How to get a consistent-looking film out of it')
                sr.italic, sr.font.size = True, Pt(11)
                sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                first_h1 = False
            else:
                h = doc.add_heading(level=min(lvl, 3))
                r = h.add_run(txt)
                r.font.color.rgb = ACCENT
                r.font.size = Pt({1: 20, 2: 15, 3: 12}.get(lvl, 11))
            i += 1
            continue

        if ln.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1]):
            header = [c.strip() for c in ln.strip('|').split('|')]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Light Grid Accent 1'
            for c, txt in zip(t.rows[0].cells, header):
                c.text = ''
                rich(c.paragraphs[0], txt)
                for rr in c.paragraphs[0].runs:
                    rr.bold = True
            for row in rows:
                cells = t.add_row().cells
                for c, txt in zip(cells, row + [''] * (len(header) - len(row))):
                    c.text = ''
                    rich(c.paragraphs[0], txt)
            doc.add_paragraph()
            continue

        if re.match(r'^[-*]\s+', ln):
            rich(doc.add_paragraph(style='List Bullet'), re.sub(r'^[-*]\s+', '', ln))
            i += 1
            continue

        if re.match(r'^\d+\.\s+', ln):
            rich(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s+', '', ln))
            i += 1
            continue

        if not ln.strip():
            i += 1
            continue

        buf = [ln]
        i += 1
        while (i < len(lines) and lines[i].strip()
               and not re.match(r'^(#{1,4}\s|[-*]\s|\d+\.\s|\|)', lines[i])
               and lines[i].strip() != '---'):
            buf.append(lines[i])
            i += 1
        rich(doc.add_paragraph(), ' '.join(x.strip() for x in buf))

    doc.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
